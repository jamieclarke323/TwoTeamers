from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from app.data_loader import _shared_players, _stats_players, load_club_players, load_featured_quizzes
from app.matching import MatchResult, match_guess

ROOT = Path(__file__).resolve().parents[2]


def _workbook(tmp_path, teams):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for club, rows in teams.items():
        seasons = workbook.create_sheet(club)
        seasons.append(["Rk", "Team", "Season"])
        seasons.append(["Count", "Player List"])
        names = list(dict.fromkeys(row[0] for row in rows))
        seasons.append([1, club, "2020-2021", len(names), *names])
        stats = workbook.create_sheet(f"{club}Stats")
        stats.append(["Player", "Yrs", "MP", "Min", "Gls", "Ast", "Nationality", "PlayerId"])
        for name, player_id in rows:
            stats.append([name, "2020-2021", 10, 900, 2, 1, "Brazil", player_id])
    path = tmp_path / "players.xlsx"
    workbook.save(path)
    workbook.close()
    return load_club_players(path, set(teams))


def test_different_ids_do_not_match_despite_identical_names(tmp_path):
    rosters = _workbook(tmp_path, {
        "Home": [("Marcelo", "603116a7")],
        "Away": [("Marcelo", "43314265")],
    })
    assert _shared_players(rosters, "Home", "Away") == []


def test_same_id_matches_different_display_names_and_preserves_stats(tmp_path):
    rosters = _workbook(tmp_path, {
        "Home": [("Player Name", "00123abc")],
        "Away": [("Different Name", "00123abc")],
    })
    players = _shared_players(rosters, "Home", "Away")
    assert len(players) == 1
    assert players[0].home.player_id == players[0].away.player_id == "00123abc"
    assert players[0].name == "Player Name"
    assert players[0].home.apps == players[0].away.apps == 10
    assert players[0].home.years == "2020-2021"
    assert players[0].clue_value("nationality", "Home", "Away") == "🇧🇷 Brazil"
    assert match_guess("Different Name", players, set()).result == MatchResult.EXACT


def test_same_named_players_at_one_club_keep_separate_ids(tmp_path):
    rosters = _workbook(tmp_path, {
        "Home": [("Same Name", "first"), ("Same Name", "second")],
        "Away": [("Same Name", "second")],
    })
    assert len(rosters["Home"]) == 2
    players = _shared_players(rosters, "Home", "Away")
    assert len(players) == 1
    assert players[0].home.player_id == players[0].away.player_id == "second"


def test_both_same_named_players_remain_separate_answers_when_shared(tmp_path):
    rosters = _workbook(tmp_path, {
        "Home": [("Same Name", "second"), ("Same Name", "first")],
        "Away": [("Same Name", "first"), ("Same Name", "second")],
    })
    players = _shared_players(rosters, "Home", "Away")
    assert [player.home.player_id for player in players] == ["first", "second"]
    assert [player.away.player_id for player in players] == ["first", "second"]
    assert match_guess("Same Name", players, {0}).player_index == 1


def test_missing_ids_only_match_other_missing_ids(tmp_path):
    rosters = _workbook(tmp_path, {
        "Home": [("Legacy Player", None), ("Known Player", "known")],
        "Away": [("Legacy Player", None), ("Known Player", None)],
    })
    players = _shared_players(rosters, "Home", "Away")
    assert [player.name for player in players] == ["Legacy Player"]


def test_stats_read_id_from_header_not_nationality_column():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Player", "Yrs", "MP", "Min", "Gls", "Ast", "PlayerId"])
    sheet.append(["Marcelo", "2020-2021", 10, 900, 2, 1, "603116a7"])
    records, ambiguous = _stats_players(sheet)
    assert not ambiguous
    player = next(iter(records.values()))
    assert player.player_id == "603116a7"
    assert player.nationality == ""
    workbook.close()


def test_stats_ignore_duplicate_headers_and_data_after_column_h():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([
        "Player", "Years", "Apps", "Min", "Goals", "Ast", "Nationality", "PlayerId",
        "Player", "PlayerId", "Nationality",
    ])
    sheet.append([
        "Lee Grant", "2018-2021", 4, 321, 0, 0, "England", "9dab6e3b",
        "David de Gea", "7ba6d84e", "Spain",
    ])
    records, ambiguous = _stats_players(sheet)
    assert not ambiguous
    assert len(records) == 1
    player = records["id:9dab6e3b"]
    assert player.name == "Lee Grant"
    assert player.nationality == "England"
    assert player.player_id == "9dab6e3b"
    workbook.close()


def test_manchester_united_atletico_answers_have_correct_names_and_ids():
    rosters = load_club_players(
        ROOT / "Two Teamers - PremierLeague.xlsx", {"Manchester United"}
    )
    rosters.update(load_club_players(
        ROOT / "Two Teamers - Champions League.xlsx", {"Atletico Madrid"}
    ))
    expected = {
        "Matheus Cunha": "dc62b55d",
        "Radamel Falcao": "66116290",
        "Diego Forlán": "d87d0dfc",
        "David de Gea": "7ba6d84e",
        "Memphis": "8f696594",
        "Sergio Reguilón": "3353737a",
    }
    grant = next(
        player for player in rosters["Manchester United"].values() if player.name == "Lee Grant"
    )
    assert grant.player_id == "9dab6e3b"
    for home, away in (
        ("Manchester United", "Atletico Madrid"),
        ("Atletico Madrid", "Manchester United"),
    ):
        players = _shared_players(rosters, home, away)
        assert {player.name: player.home.player_id for player in players} == expected
        assert len(players) == len(expected)
        for player in players:
            assert player.home.player_id == player.away.player_id == expected[player.name]
            assert player.home.name == player.away.name == player.name
        assert match_guess("Lee Grant", players, set()).result == MatchResult.WRONG
        assert match_guess("Eric Djemba-Djemba", players, set()).result == MatchResult.WRONG


def test_identified_pre_cutoff_namesake_is_not_added(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Home"
    sheet.append(["Rk", "Team", "Season"])
    sheet.append(["Count", "Player List"])
    sheet.append([1, "Home", "2020-2021", 1, "Same Name"])
    stats = workbook.create_sheet("HomeStats")
    stats.append(["Player", "Yrs", "MP", "Min", "Gls", "Ast", "Nationality", "PlayerId"])
    stats.append(["Same Name", "1980-1985", 100, None, 2, None, None, "old"])
    stats.append(["Same Name", "2020-2021", 10, None, 1, None, None, "new"])
    path = tmp_path / "cutoff.xlsx"
    workbook.save(path)
    workbook.close()
    players = load_club_players(path, {"Home"})["Home"]
    assert [player.player_id for player in players.values()] == ["new"]


def test_ambiguous_id_does_not_fall_back_to_name_matching(tmp_path):
    rosters = _workbook(tmp_path, {
        "Home": [("Same Name", "duplicate"), ("Same Name", "duplicate")],
        "Away": [("Same Name", None)],
    })
    assert not rosters["Home"]
    assert _shared_players(rosters, "Home", "Away") == []


def test_real_madrid_and_birmingham_marcelos_are_different_people(tmp_path):
    rosters = load_club_players(
        ROOT / "Two Teamers - Champions League.xlsx", {"Real Madrid"}
    )
    rosters.update(load_club_players(
        ROOT / "Two Teamers - Championship.xlsx", {"Birmingham City"}
    ))
    madrid = next(player for player in rosters["Real Madrid"].values() if player.name == "Marcelo")
    birmingham = next(
        player for player in rosters["Birmingham City"].values() if player.name == "Marcelo"
    )
    assert madrid.player_id == "603116a7"
    assert birmingham.player_id == "43314265"
    assert all(
        player.name != "Marcelo"
        for player in _shared_players(rosters, "Real Madrid", "Birmingham City")
    )
    assert [
        player.name for player in _shared_players(rosters, "Real Madrid", "Birmingham City")
    ] == ["Jude Bellingham"]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Featured fixture list"
    sheet.append(["Date", "Kick off", "Matchday", "Home", "Away", "City", "Stadium"])
    sheet.append([
        datetime(2026, 10, 22), "WILDCARD", None,
        "Real Madrid", "Birmingham City", None, None,
    ])
    path = tmp_path / "featured.xlsx"
    workbook.save(path)
    workbook.close()
    quizzes = load_featured_quizzes(path, rosters, set(rosters))
    assert len(quizzes) == 1
    assert [player.name for player in quizzes[0].players] == ["Jude Bellingham"]
