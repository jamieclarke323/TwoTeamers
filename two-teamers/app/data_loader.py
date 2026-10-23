from __future__ import annotations

import csv
import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime, time, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl

QUIZ_TIMEZONE = ZoneInfo("Europe/London")
UNLOCK_HOUR = 6
FEATURED_UNLOCK = time(0, 1)
PLAYABLE_HOURS = 100
MIN_ELIGIBLE_SEASON_START = 1999

LEAGUE_FILES = {
    "premier-league": "2627PremierLeagueFixtures.csv",
    "championship": "2627ChampionshipFixtures.csv",
    "league-one": "2627LeagueOneFixtures.csv",
    "league-two": "2627LeagueTwoFixtures.csv",
}
LEAGUE_NAMES = {
    "premier-league": "Premier League",
    "championship": "Championship",
    "league-one": "League One",
    "league-two": "League Two",
}
PLAYER_WORKBOOKS = {
    "premier-league": "Two Teamers - PremierLeague.xlsx",
    "championship": "Two Teamers - Championship.xlsx",
    "league-one": "Two Teamers - League One.xlsx",
    "league-two": "Two Teamers - League Two.xlsx",
}
LEAGUE_PREFIXES = {
    "premier-league": "pl",
    "championship": "ch",
    "league-one": "l1",
    "league-two": "l2",
}
CLUB_ALIASES = {
    "Brighton": "Brighton & Hove Albion",
    "Coventry": "Coventry City",
    "Hull": "Hull City",
    "Ipswich": "Ipswich Town",
    "Leeds": "Leeds United",
    "Man City": "Manchester City",
    "Man Utd": "Manchester United",
    "Newcastle": "Newcastle United",
    "Nott'm Forest": "Nottingham Forest",
    "Spurs": "Tottenham Hotspur",
}
FEATURED_WORKBOOK = "Featured fixtures.xlsx"
FEATURED_SHEET = "Featured fixture list"
FEATURED_LEAGUE_ID = "featured"
CHAMPIONS_LEAGUE_WORKBOOK = "Two Teamers - Champions League.xlsx"
FEATURED_CLUB_ALIASES = {
    "Ason Villa": "Aston Villa",
    "Bolton": "Bolton Wanderers",
    "Brighton and Hove Albion": "Brighton & Hove Albion",
    "Cardiff": "Cardiff City",
    "Charlton": "Charlton Athletic",
    "Leicester": "Leicester City",
    "Sheffield Utd": "Sheffield United",
    "Tottenham": "Tottenham Hotspur",
    "West Ham": "West Ham United",
}
WORKBOOK_SHEET_ALIASES = {
    "Brighton & Hove Albion": "Brighton",
    "Coventry City": "Coventry",
    "Norwich City": "Norwich",
    "Peterborough United": "Peterborough",
    "Chesterfield": "Chesterfield Town",
    "Tranmere Rovers": "Tranmere",
}

_EXTRA_TRANSLITERATIONS = str.maketrans(
    {"ð": "d", "Ð": "d", "þ": "th", "Þ": "th", "æ": "ae", "Æ": "ae", "ø": "o", "Ø": "o", "ł": "l", "Ł": "l"}
)


def slugify(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def normalise_name(value: str) -> str:
    value = value.translate(_EXTRA_TRANSLITERATIONS)
    value = unicodedata.normalize("NFKD", value)
    value = "".join(char for char in value if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]", "", value.lower())


def clean_value(value) -> str | int:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


_NATIONALITY_FLAG_CODES = {
    "Albania": "AL", "Algeria": "DZ", "Angola": "AO", "Argentina": "AR",
    "Armenia": "AM", "Australia": "AU", "Austria": "AT", "Barbados": "BB",
    "Belarus": "BY", "Belgium": "BE", "Benin": "BJ", "Bermuda": "BM",
    "Bosnia and Herzegovina": "BA", "Brazil": "BR", "Bulgaria": "BG",
    "Burkina Faso": "BF", "Burundi": "BI", "Cameroon": "CM", "Canada": "CA",
    "Chile": "CL", "China": "CN", "Colombia": "CO", "Costa Rica": "CR",
    "Croatia": "HR", "Curaçao": "CW", "Cyprus": "CY", "Czech Republic": "CZ",
    "DR Congo": "CD", "Democratic Republic of the Congo": "CD", "Denmark": "DK",
    "Ecuador": "EC", "Egypt": "EG", "Equatorial Guinea": "GQ", "Estonia": "EE",
    "Faroe Islands": "FO", "Finland": "FI", "France": "FR", "Gabon": "GA",
    "Gambia": "GM", "Georgia": "GE", "Germany": "DE", "Ghana": "GH",
    "Greece": "GR", "Grenada": "GD", "Guinea": "GN", "Guinea-Bissau": "GW",
    "Guyana": "GY", "Haiti": "HT", "Honduras": "HN", "Hungary": "HU",
    "Iceland": "IS", "Indonesia": "ID", "Iran": "IR", "Iraq": "IQ",
    "Ireland": "IE", "Republic of Ireland": "IE", "Israel": "IL", "Italy": "IT",
    "Ivory Coast": "CI", "Jamaica": "JM", "Japan": "JP", "Jersey": "JE",
    "Kenya": "KE", "Latvia": "LV", "Liberia": "LR",
    "Lithuania": "LT", "Mali": "ML", "Malta": "MT", "Martinique": "MQ",
    "Mauritania": "MR", "Mauritius": "MU", "Mexico": "MX", "Montserrat": "MS",
    "Morocco": "MA", "Netherlands": "NL", "New Zealand": "NZ", "Nigeria": "NG",
    "Norway": "NO", "Oman": "OM", "Pakistan": "PK", "Paraguay": "PY",
    "Peru": "PE", "Poland": "PL", "Portugal": "PT", "Republic of the Congo": "CG",
    "Romania": "RO", "Russia": "RU", "Saint Kitts and Nevis": "KN",
    "Saint Lucia": "LC", "Senegal": "SN", "Serbia": "RS", "Sierra Leone": "SL",
    "Slovakia": "SK", "Slovenia": "SI", "South Africa": "ZA", "South Korea": "KR",
    "Spain": "ES", "Sweden": "SE", "Switzerland": "CH", "Tanzania": "TZ",
    "Togo": "TG", "Trinidad and Tobago": "TT", "Tunisia": "TN", "Turkey": "TR",
    "Uganda": "UG", "Ukraine": "UA", "United States": "US", "Uruguay": "UY",
    "Venezuela": "VE", "Zambia": "ZM", "Zimbabwe": "ZW",
}
_UK_NATIONALITY_FLAGS = {
    "England": "\U0001f3f4\U000e0067\U000e0062\U000e0065\U000e006e\U000e0067\U000e007f",
    "Scotland": "\U0001f3f4\U000e0067\U000e0062\U000e0073\U000e0063\U000e0074\U000e007f",
    "Wales": "\U0001f3f4\U000e0067\U000e0062\U000e0077\U000e006c\U000e0073\U000e007f",
}


def nationality_flag(nationality: str) -> str:
    """Return a country flag emoji, or the original nationality if unavailable."""
    if nationality in _UK_NATIONALITY_FLAGS:
        return _UK_NATIONALITY_FLAGS[nationality]
    code = _NATIONALITY_FLAG_CODES.get(nationality)
    if not code:
        return nationality
    return "".join(chr(0x1F1E6 + ord(letter) - ord("A")) for letter in code)


def nationality_display(nationality: str) -> str:
    """Return a nationality with its flag emoji when one is available."""
    flag = nationality_flag(nationality)
    return f"{flag} {nationality}" if flag != nationality else nationality


def _nationality_identity(nationality: str) -> tuple[str, str]:
    if nationality in _NATIONALITY_FLAG_CODES:
        return "country-code", _NATIONALITY_FLAG_CODES[nationality]
    if nationality in _UK_NATIONALITY_FLAGS:
        return "subdivision", nationality
    return "name", nationality.casefold()


@dataclass(frozen=True)
class ClubPlayer:
    name: str
    years: str
    apps: str | int
    goals: str | int
    nationality: str = ""
    player_id: str = ""

    @property
    def identity(self) -> str:
        return f"id:{self.player_id}" if self.player_id else normalise_name(self.name)


@dataclass
class QuizPlayer:
    name: str
    home: ClubPlayer
    away: ClubPlayer
    surname: str = field(init=False)
    acceptable_answers: list[str] = field(init=False)
    initials: str = field(init=False)

    def __post_init__(self) -> None:
        parts = self.name.split()
        self.surname = parts[-1] if parts else self.name
        candidates = []
        for name in dict.fromkeys((self.name, self.home.name, self.away.name)):
            name_parts = name.split()
            candidates.append(name)
            if len(name_parts) >= 3:
                candidates.append(" ".join(name_parts[-2:]))
            if name_parts:
                candidates.append(name_parts[-1])
        self.acceptable_answers = list(dict.fromkeys(candidates))
        self.initials = "".join(f"{part[0].upper()}." for part in parts if part)

    def clue_value(self, clue_type: str, home_team: str, away_team: str) -> str:
        if clue_type == "years":
            return f"{home_team}: {self.home.years} / {away_team}: {self.away.years}"
        if clue_type == "initials":
            return self.initials
        if clue_type == "nationality":
            nationalities = {}
            for nationality in (self.home.nationality, self.away.nationality):
                nationalities.setdefault(_nationality_identity(nationality), nationality)
            return " / ".join(nationality_display(value) for value in nationalities.values())
        raise ValueError(f"Unknown clue type: {clue_type}")

    @property
    def has_nationality(self) -> bool:
        return bool(self.home.nationality and self.away.nationality)


@dataclass
class Quiz:
    quiz_id: str
    league_id: str
    match_number: int
    matchweek: int
    fixture_date: datetime
    venue: str
    home_team: str
    away_team: str
    players: list[QuizPlayer]
    wildcard: bool = False

    @property
    def is_featured(self) -> bool:
        return self.league_id == FEATURED_LEAGUE_ID

    @property
    def unlock_at(self) -> datetime:
        unlock_time = FEATURED_UNLOCK if self.is_featured else time(UNLOCK_HOUR)
        return datetime.combine(self.fixture_date.date(), unlock_time, QUIZ_TIMEZONE)

    @property
    def lock_at(self) -> datetime:
        return self.unlock_at + timedelta(hours=PLAYABLE_HOURS)

    @property
    def total_players(self) -> int:
        return len(self.players)

    def availability(self, now: datetime | None = None) -> str:
        now = now or datetime.now(QUIZ_TIMEZONE)
        if now.tzinfo is None:
            now = now.replace(tzinfo=QUIZ_TIMEZONE)
        if now < self.unlock_at:
            return "upcoming"
        if now >= self.lock_at:
            return "expired"
        return "playable"


@dataclass(frozen=True)
class League:
    league_id: str
    name: str
    clubs: tuple[str, ...]
    enabled: bool


def _club_name(csv_name: str) -> str:
    return CLUB_ALIASES.get(csv_name, csv_name)


def _unique_season_players(sheet) -> dict[str, tuple[str, set[str], bool]]:
    players: dict[str, set[str]] = defaultdict(set)
    seasons: dict[str, set[str]] = defaultdict(set)
    current_season_players: set[str] = set()
    found_current_season = False
    for raw_row in sheet.iter_rows(min_row=3, values_only=True):
        row = tuple(raw_row) + (None,) * max(0, 5 - len(raw_row))
        try:
            player_count = int(row[3] or 0)
        except (TypeError, ValueError):
            continue
        season = str(row[2] or "").strip()
        is_current_season_row = not found_current_season and bool(season) and player_count > 0
        if is_current_season_row:
            found_current_season = True
        for value in row[4 : 4 + player_count]:
            name = str(value or "").strip()
            if name:
                key = normalise_name(name)
                players[key].add(name)
                if season:
                    seasons[key].add(season)
                if is_current_season_row:
                    current_season_players.add(key)
    return {
        key: (next(iter(names)), seasons.get(key, set()), key in current_season_players)
        for key, names in players.items()
        if len(names) == 1
    }


def _stats_players(sheet) -> tuple[dict[str, ClubPlayer], set[str]]:
    records: dict[str, list[ClubPlayer]] = defaultdict(list)
    nationality_column: int | None = 6
    player_id_column: int | None = None
    found_header = False
    for raw_row in sheet.iter_rows(max_col=8, values_only=True):
        row = tuple(raw_row) + (None,) * max(0, 7 - len(raw_row))
        name = str(row[0] or "").strip()
        years = row[1]
        if years in ("Years", "Yrs"):
            if not found_header:
                columns = {str(value or "").strip(): index for index, value in enumerate(row)}
                nationality_column = columns.get("Nationality")
                player_id_column = columns.get("PlayerId")
                found_header = True
            continue
        if not name or years is None:
            continue
        nationality = (
            str(row[nationality_column] or "").strip()
            if nationality_column is not None and nationality_column < len(row)
            else ""
        )
        if nationality in ("-", "Unverified", "Footballing nationality"):
            nationality = ""
        player_id = (
            str(clean_value(row[player_id_column])).strip()
            if player_id_column is not None and player_id_column < len(row)
            else ""
        )
        player = ClubPlayer(
            name=name,
            years=str(clean_value(years)),
            apps=clean_value(row[2]),
            goals=clean_value(row[4]),
            nationality=nationality,
            player_id=player_id,
        )
        records[player.identity].append(player)
    unique = {}
    ambiguous = set()
    for key, values in records.items():
        if len(values) == 1:
            unique[key] = values[0]
            continue
        eligible = [record for record in values if _played_in_eligible_season(record.years)]
        if len(eligible) == 1:
            unique[key] = eligible[0]
        else:
            ambiguous.add(key)
            ambiguous.update(normalise_name(record.name) for record in values)
    return unique, ambiguous


def _played_in_eligible_season(years: str) -> bool:
    """Return true if a season year or year range includes 1999-2000 or later."""
    season_years = [int(value) for value in re.findall(r"\d{4}", years)]
    if not season_years:
        return False
    # A single year such as 1999 means the 1999-2000 season. Ranges are
    # calendar-year bounds, so 1998-1999 is pre-cutoff while 1998-2001 spans
    # at least one eligible season.
    if len(season_years) == 1:
        return season_years[0] >= MIN_ELIGIBLE_SEASON_START
    return max(season_years) >= MIN_ELIGIBLE_SEASON_START + 1


def load_club_players(xlsx_path: str | Path, clubs: set[str]) -> dict[str, dict[str, ClubPlayer]]:
    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    rosters: dict[str, dict[str, ClubPlayer]] = {}
    try:
        for club in sorted(clubs):
            sheet_name = WORKBOOK_SHEET_ALIASES.get(club, club)
            if sheet_name not in workbook.sheetnames:
                continue
            stats_sheet_name = next(
                (name for name in (f"{sheet_name}Stats", f"{sheet_name} Stats")
                 if name in workbook.sheetnames),
                None,
            )
            season_players = _unique_season_players(workbook[sheet_name])
            stats_players, ambiguous_stats = (
                _stats_players(workbook[stats_sheet_name])
                if stats_sheet_name is not None
                else ({}, set())
            )
            stats_by_name: dict[str, list[ClubPlayer]] = defaultdict(list)
            for player in stats_players.values():
                stats_by_name[normalise_name(player.name)].append(player)
            roster: dict[str, ClubPlayer] = {}
            for name_key, (name, seasons, _) in season_players.items():
                if not any(_played_in_eligible_season(season) for season in seasons):
                    continue
                candidates = stats_by_name.get(name_key, [])
                if not candidates:
                    if name_key not in ambiguous_stats:
                        roster[name_key] = ClubPlayer(name, "", "", "")
                    continue
                for stats in candidates:
                    # Season lists contain names, so the stats years also have to
                    # rule out a pre-cutoff namesake with a different ID.
                    if stats.player_id and not _played_in_eligible_season(stats.years):
                        continue
                    roster[stats.identity] = ClubPlayer(
                        name=name,
                        years=stats.years,
                        apps=stats.apps,
                        goals=stats.goals,
                        nationality=stats.nationality,
                        player_id=stats.player_id,
                    )
            rosters[club] = roster
    finally:
        workbook.close()
    return rosters


def load_fixture_rows(data_dir: str | Path, league_id: str) -> list[dict[str, str]]:
    path = Path(data_dir) / LEAGUE_FILES[league_id]
    with path.open(encoding="utf-8-sig", newline="") as fixture_file:
        return list(csv.DictReader(fixture_file))


def load_all_clubs(data_dir: str | Path) -> dict[str, list[str]]:
    clubs_by_league: dict[str, list[str]] = {}
    for league_id in LEAGUE_FILES:
        rows = load_fixture_rows(data_dir, league_id)
        clubs = {_club_name(row[side].strip()) for row in rows for side in ("Home Team", "Away Team")}
        clubs_by_league[league_id] = sorted(clubs)
    return clubs_by_league


def _shared_players(rosters, home_team: str, away_team: str) -> list[QuizPlayer]:
    home_roster = rosters.get(home_team, {})
    away_roster = rosters.get(away_team, {})
    shared_names = sorted(
        home_roster.keys() & away_roster.keys(),
        key=lambda key: (
            normalise_name(home_roster[key].name.split()[-1]),
            normalise_name(home_roster[key].name),
            key,
        ),
    )
    return [QuizPlayer(home_roster[key].name, home_roster[key], away_roster[key]) for key in shared_names]


def _featured_club_name(value) -> str:
    name = str(value or "").strip()
    return _club_name(FEATURED_CLUB_ALIASES.get(name, name))


def load_featured_quizzes(xlsx_path: str | Path, rosters, known_clubs: set[str]) -> list[Quiz]:
    """Build one quiz per dated row; rows without two recognised clubs are skipped."""
    path = Path(xlsx_path)
    if not path.is_file():
        return []
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if FEATURED_SHEET not in workbook.sheetnames:
        return []
    quizzes: list[Quiz] = []
    for row in workbook[FEATURED_SHEET].iter_rows(min_row=2, values_only=True):
        row = tuple(row) + (None,) * max(0, 7 - len(row))
        day, kick_off, _matchday, home_value, away_value, _city, stadium = row[:7]
        if not isinstance(day, datetime):
            continue
        home_team, away_team = _featured_club_name(home_value), _featured_club_name(away_value)
        if home_team not in known_clubs or away_team not in known_clubs:
            continue
        wildcard = str(kick_off or "").strip().upper() == "WILDCARD"
        fixture_date = day.replace(tzinfo=QUIZ_TIMEZONE)
        if not wildcard and re.fullmatch(r"\d{1,2}:\d{2}", str(kick_off or "").strip()):
            hour, minute = map(int, str(kick_off).strip().split(":"))
            fixture_date = fixture_date.replace(hour=hour, minute=minute)
        quizzes.append(
            Quiz(
                quiz_id=f"featured-{day.date().isoformat()}-{slugify(home_team)}-{slugify(away_team)}",
                league_id=FEATURED_LEAGUE_ID,
                match_number=0,
                matchweek=0,
                fixture_date=fixture_date,
                venue=str(stadium or "").strip(),
                home_team=home_team,
                away_team=away_team,
                players=_shared_players(rosters, home_team, away_team),
                wildcard=wildcard,
            )
        )
    quizzes.sort(key=lambda quiz: quiz.fixture_date)
    return quizzes


def load_champions_league_rosters(xlsx_path: str | Path) -> dict[str, dict[str, ClubPlayer]]:
    path = Path(xlsx_path)
    if not path.is_file():
        return {}
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        clubs = {name for name in workbook.sheetnames if not name.endswith("Stats") and name != "TeamIDs"}
    finally:
        workbook.close()
    return load_club_players(path, clubs)


def load_quizzes(data_dir: str | Path, _legacy_xlsx_path: str | Path | None = None) -> tuple[list[League], list[Quiz]]:
    leagues, quizzes, *_ = load_game_data(data_dir)
    return leagues, quizzes


def load_game_data(
    data_dir: str | Path,
    featured_path: str | Path | None = None,
    cl_path: str | Path | None = None,
) -> tuple[list[League], list[Quiz], list[Quiz], dict[str, dict[str, ClubPlayer]]]:
    data_dir = Path(data_dir)
    clubs_by_league = load_all_clubs(data_dir)

    quizzes: list[Quiz] = []
    all_rosters: dict[str, dict[str, ClubPlayer]] = {}
    for league_id, workbook_name in PLAYER_WORKBOOKS.items():
        clubs = set(clubs_by_league[league_id])
        rosters = load_club_players(data_dir / workbook_name, clubs)
        all_rosters.update(rosters)
        prefix = LEAGUE_PREFIXES[league_id]
        for row in load_fixture_rows(data_dir, league_id):
            home_team = _club_name(row["Home Team"].strip())
            away_team = _club_name(row["Away Team"].strip())
            players = _shared_players(rosters, home_team, away_team)
            match_number = int(row["Match Number"])
            quizzes.append(
                Quiz(
                    quiz_id=f"{prefix}-{match_number}-{slugify(home_team)}-{slugify(away_team)}",
                    league_id=league_id,
                    match_number=match_number,
                    matchweek=int(row["Round Number"]),
                    fixture_date=datetime.strptime(row["Date"], "%d/%m/%Y %H:%M").replace(tzinfo=QUIZ_TIMEZONE),
                    venue=row["Location"].strip(),
                    home_team=home_team,
                    away_team=away_team,
                    players=players,
                )
            )

    quizzes.sort(key=lambda quiz: (quiz.fixture_date, quiz.match_number))
    leagues = [
        League(league_id, LEAGUE_NAMES[league_id], tuple(clubs_by_league[league_id]), league_id in PLAYER_WORKBOOKS)
        for league_id in LEAGUE_FILES
    ]
    cl_file = Path(cl_path) if cl_path else data_dir / CHAMPIONS_LEAGUE_WORKBOOK
    cl_rosters = load_champions_league_rosters(cl_file)
    all_rosters.update(cl_rosters)

    known_clubs = {club for clubs in clubs_by_league.values() for club in clubs} | set(cl_rosters.keys())
    featured = load_featured_quizzes(featured_path, all_rosters, known_clubs) if featured_path else []
    return leagues, quizzes, featured, all_rosters
